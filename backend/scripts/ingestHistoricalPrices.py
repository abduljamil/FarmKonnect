"""
One-shot ingestion: load 17 years of AMIS Pakistan commodity prices from the
local xlsx workbooks at `C:\\Users\\ahtsh\\OneDrive\\Desktop\\data collection\\
complete data\\` into the live `commodityprices` collection on Atlas.

Idempotent: safe to interrupt and re-run. Uses the existing unique compound
index {commodity, variety, city, date, priceType} as the upsert filter, so
nothing is duplicated.

Format choices match the live production scraper:
  - commodity = full sheet/workbook name (rice variants are NOT split into
    commodity+variety; they live as their own commodity strings)
  - variety = None
  - city = CamelCase concatenated (BahawalNagar, RahimYarKhan, BahawalPur)
  - priceType = "FQP"
  - unit = "Rs/100Kg"   (what AMIS CommodityChart exports)
  - date = native datetime (not string)

Throttled for Atlas M0 free tier: batches of 500 with 200 ms pause.

Usage:
    # Read MONGODB_URI from env, NEVER hard-code it
    set MONGODB_URI=<atlas-uri>
    python backend/scripts/ingestHistoricalPrices.py --dry-run
    python backend/scripts/ingestHistoricalPrices.py
"""

from __future__ import annotations

import argparse
import os
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

try:
    import openpyxl
    from pymongo import MongoClient, UpdateOne
    from pymongo.errors import BulkWriteError
except ImportError:
    print("Install deps:  pip install openpyxl pymongo")
    sys.exit(1)


# -----------------------------------------------------------------------------
# Configuration
# -----------------------------------------------------------------------------
DATA_DIR = Path(r"C:\Users\ahtsh\OneDrive\Desktop\data collection\complete data")

# (workbook filename, sheet name -> commodity string)
WORKBOOKS = {
    "Wheat_Prices_AMIS.xlsx":          {"Wheat": "Wheat"},
    "Maize_Prices_AMIS.xlsx":          {"Maize": "Maize"},
    "Sugar_Prices_AMIS.xlsx":          {"Sugar": "Sugar"},
    "Cotton_Phutti_Prices_AMIS.xlsx":  {"Seed Cotton (Phutti)": "Seed Cotton (Phutti)"},
    "Rice_All_Varieties_Prices_AMIS.xlsx": {
        "Rice Basmati Super (New)":  "Rice Basmati Super (New)",
        "Rice (IRRI)":               "Rice (IRRI)",
        "Paddy Basmati":             "Paddy Basmati",
        "Paddy (IRRI)":              "Paddy (IRRI)",
        "Rice Basmati Super (Old)":  "Rice Basmati Super (Old)",
        "Rice Basmati (385)":        "Rice Basmati (385)",
        "Paddy Kainat":              "Paddy Kainat",
        "Rice Kainat (New)":         "Rice Kainat (New)",
    },
}

# xlsx-side city label -> Mongo-side CamelCase city
CITY_MAP = {
    "Lahore":         "Lahore",
    "Faisalabad":     "Faisalabad",
    "Gujranwala":     "Gujranwala",
    "Okara":          "Okara",
    "Sargodha":       "Sargodha",
    "Rawalpindi":     "Rawalpindi",
    "Multan":         "Multan",
    "Rahim Yar Khan": "RahimYarKhan",
    "Layyah":         "Layyah",
    "Bahawalpur":     "BahawalPur",
    "Bahawalnagar":   "BahawalNagar",
    "Chichawatni":    "Chichawatni",
    "Sialkot":        "Sialkot",
    "Jhang":          "Jhang",
}

MONTH_TO_NUM = {
    "January": 1, "February": 2, "March": 3, "April": 4, "May": 5, "June": 6,
    "July": 7, "August": 8, "September": 9, "October": 10, "November": 11, "December": 12,
}

# Drop dead series before ingest -- they have so few rows they're noise.
DROP_SERIES = {
    ("Maize",                   "BahawalNagar"),    # 47 rows over 10 years
    ("Seed Cotton (Phutti)",    "Faisalabad"),      # 5 rows
    ("Seed Cotton (Phutti)",    "Lahore"),          # 2 rows
    ("Seed Cotton (Phutti)",    "Rawalpindi"),      # 1 row
}

BATCH_SIZE = 500
SLEEP_BETWEEN_BATCHES_S = 0.20


# -----------------------------------------------------------------------------
# Loaders
# -----------------------------------------------------------------------------
def load_workbook_records(xlsx_path: Path, sheet_to_commodity: dict[str, str]):
    """Yield (commodity, city, date, price) tuples from one workbook."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    for sheet_name, commodity in sheet_to_commodity.items():
        if sheet_name not in wb.sheetnames:
            print(f"  WARN: sheet '{sheet_name}' not in {xlsx_path.name}, skipping")
            continue
        ws = wb[sheet_name]
        header = None
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                header = row
                continue
            try:
                city_raw, year, month_name, day, value = row[0], row[1], row[2], row[3], row[4]
                if city_raw is None or value is None:
                    continue
                city = CITY_MAP.get(str(city_raw).strip())
                if city is None:
                    print(f"  WARN: unknown city '{city_raw}' (skipping)")
                    continue
                if (commodity, city) in DROP_SERIES:
                    continue
                m = MONTH_TO_NUM.get(str(month_name).strip())
                if m is None:
                    continue
                dt = datetime(int(year), m, int(day))
                price = float(value)
                yield (commodity, city, dt, price)
            except (TypeError, ValueError):
                continue
    wb.close()


def build_doc(commodity: str, city: str, date: datetime, price: float, now: datetime) -> dict:
    return {
        "commodity":   commodity,
        "variety":     None,
        "city":        city,
        "price":       price,
        "priceType":   "FQP",
        "unit":        "Rs/100Kg",
        "date":        date,
        "timestamp":   now,
        "lastUpdated": now,
    }


# -----------------------------------------------------------------------------
# Main
# -----------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true",
                    help="Read xlsx and build docs, but do not write to Mongo.")
    ap.add_argument("--db", default="FarmKonnect")
    ap.add_argument("--collection", default="commodityprices")
    args = ap.parse_args()

    uri = os.environ.get("MONGODB_URI") or os.environ.get("MONGO_URI")
    if not args.dry_run and not uri:
        print("ERROR: MONGODB_URI not set in environment")
        sys.exit(2)

    print("=" * 70)
    print("AMIS Historical Price Ingestion")
    print("=" * 70)
    print(f"  data dir   : {DATA_DIR}")
    print(f"  dry-run    : {args.dry_run}")
    print(f"  target db  : {args.db}.{args.collection}")
    print(f"  batch size : {BATCH_SIZE}")
    print()

    # Connect (skip if dry-run with no URI)
    coll = None
    if not args.dry_run:
        client = MongoClient(uri)
        coll = client[args.db][args.collection]
        before = coll.count_documents({})
        print(f"  current docs in collection: {before:,}")
        print()

    # Load all workbooks
    now_ts = datetime.now(timezone.utc)
    total_built = 0
    total_inserted = 0
    total_updated = 0
    total_errors = 0
    per_series = {}                                # (commodity, city) -> count

    pending: list[UpdateOne] = []

    def flush(label: str = ""):
        nonlocal total_inserted, total_updated, total_errors, pending
        if not pending:
            return
        if args.dry_run:
            pending = []
            return
        try:
            res = coll.bulk_write(pending, ordered=False)
            total_inserted += res.upserted_count
            total_updated += res.modified_count
        except BulkWriteError as e:
            total_errors += len(e.details.get("writeErrors", []))
            # account whatever did succeed
            total_inserted += e.details.get("nUpserted", 0)
            total_updated  += e.details.get("nModified", 0)
            for we in e.details.get("writeErrors", [])[:3]:
                print(f"    write err: {we.get('errmsg','?')[:200]}")
        pending = []
        time.sleep(SLEEP_BETWEEN_BATCHES_S)

    for fname, sheet_map in WORKBOOKS.items():
        path = DATA_DIR / fname
        if not path.exists():
            print(f"MISSING: {path}"); continue
        print(f"\n>> {fname}")
        for commodity, city, dt, price in load_workbook_records(path, sheet_map):
            total_built += 1
            per_series[(commodity, city)] = per_series.get((commodity, city), 0) + 1
            doc = build_doc(commodity, city, dt, price, now_ts)
            pending.append(UpdateOne(
                filter={
                    "commodity": commodity, "variety": None, "city": city,
                    "date": dt, "priceType": "FQP",
                },
                update={"$set": doc},
                upsert=True,
            ))
            if len(pending) >= BATCH_SIZE:
                flush()
                if total_built % (BATCH_SIZE * 20) == 0:
                    print(f"  ... {total_built:,} rows built, "
                          f"{total_inserted:,} ins / {total_updated:,} upd so far")
    flush(label="final")

    print()
    print("=" * 70)
    print("DONE")
    print("=" * 70)
    print(f"  rows built   : {total_built:,}")
    if not args.dry_run:
        print(f"  inserted     : {total_inserted:,}")
        print(f"  updated      : {total_updated:,}")
        print(f"  errors       : {total_errors:,}")
        after = coll.count_documents({})
        print(f"  collection size now: {after:,} (was {before:,}, delta {after-before:+,})")

    print()
    print("Per-series row counts:")
    for (com, city), n in sorted(per_series.items()):
        print(f"  {com:30s}  {city:18s}  {n:6,d}")


if __name__ == "__main__":
    main()
